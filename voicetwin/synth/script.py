"""讲稿解析：把 txt / md / srt / docx 讲稿切成适合合成的句子。

讲稿里可以用的标记：
    空一行                     → 段落停顿（按你本人的段落停顿习惯）
    [停顿] / [pause]           → 段落停顿
    [停顿=1.5] / [停顿 2秒]    → 指定停顿秒数（[pause=1.5s]、<break time="800ms"/> 也可以）
中文输入法打出来的【停顿=2】【停顿】［停顿＝２］也认（括号、等号、数字是全角的也行）；
写在讲稿最前面 / 最后面的停顿秒数，加在开头 / 结尾的静音上（比自带的留白短时用自带的）。
Markdown 的标题、列表、加粗、链接会被自动清理；``` 代码块默认不朗读。
"""

from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple, Union

from voicetwin.utils.textutil import (
    CLAUSE_CHARS,
    SENT_END_CHARS,
    count_cjk,
    detect_lang,
    ensure_final_punct,
    sentence_kind,
    syllable_count,
)

PAUSE_RE = re.compile(
    r"\[(?:停顿|暂停|pause|break)(?:\s*[=:：]?\s*(\d+(?:\.\d+)?)\s*(ms|毫秒|s|秒钟|秒|分钟|min)?)?\s*\]"
    r"|<break\s+time\s*=\s*[\"']?(\d+(?:\.\d+)?)\s*(ms|s)?[\"']?\s*/?>",
    re.IGNORECASE,
)
#: 中文输入法下打 [ ] 会变成【 】（微软拼音、搜狗默认都这样），等号、数字也可能是全角的：
#: 【停顿=2】【停顿】［停顿＝２］[停顿＝2] 先统一成 [停顿=2] 再认。括号里最多 20 个字、不跨行，
#: 统一以后认不出来的（例如【暂停一下】这种小标题）原样不动
_PAUSE_LOOSE_RE = re.compile(r"[\[【［]\s*((?:停顿|暂停|pause|break)[^\]】］\n]{0,20}?)\s*[\]】］]", re.IGNORECASE)
#: 统一以后还剩下的、括号里写着停顿的标记（例如【停顿一下】【停顿=2分半】）：会被当成文字读出来，要提醒老师
_UNREAD_MARK_RE = re.compile(r"[\[【［]\s*(?:停顿|暂停)[^\]】］\n]{0,20}[\]】］]")
ABBREVIATIONS = {"mr", "mrs", "ms", "dr", "prof", "sr", "jr", "st", "vs", "etc", "e.g", "i.e", "no", "fig", "eq",
                 "approx", "dept", "inc", "ltd", "co", "u.s", "a.m", "p.m"}
EMOJI_RE = re.compile("[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F2FF]+")


@dataclass
class ScriptSegment:
    text: str                           # 送去合成的文字（已应用读音词典）
    display: str                        # 字幕显示的原文
    lang: str
    kind: str = "statement"
    pause_after: Union[str, float] = "sentence"   # clause | sentence | paragraph | 秒数
    paragraph: int = 0
    cue_start: Optional[float] = None   # 按字幕时间轴合成时使用
    cue_end: Optional[float] = None
    index: int = 0
    extra: dict = field(default_factory=dict)


# ----------------------------------------------------------------------------- 读取
SCRIPT_EXTS = (".txt", ".md", ".markdown", ".srt", ".vtt", ".docx")
UNSUPPORTED_MSG = ("不支持这种文件（.doc / .wps / .pdf）。请在 Word 或 WPS 里点「文件 → 另存为」，"
                   "类型选 .docx 或 .txt，再上传。")


def _looks_binary(text: str) -> bool:
    """控制字符超过 5%：多半是 .doc/.pdf 之类的二进制文件被改了扩展名，读出来会是一堆乱码。"""
    if not text:
        return False
    sample = text[:20000]
    bad = sum(1 for ch in sample if (ord(ch) < 32 and ch not in "\t\n\r\f") or ch == "\ufffd")
    return bad / max(len(sample), 1) > 0.05


def read_script_file(path: Path) -> Tuple[str, Optional[list]]:
    """返回 (纯文本, 字幕 cue 列表或 None)。只支持 .txt / .md / .srt / .vtt / .docx。"""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SCRIPT_EXTS:
        raise RuntimeError(UNSUPPORTED_MSG)
    if ext in (".srt", ".vtt"):
        from voicetwin.data.subtitles import parse_subtitles

        return "", parse_subtitles(path)
    if ext == ".docx":
        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise RuntimeError("读取 Word 讲稿需要：pip install python-docx（或另存为 txt）") from exc
        try:
            doc = docx.Document(str(path))
        except Exception as exc:  # 改了扩展名的 .doc / 损坏的文件
            raise RuntimeError(UNSUPPORTED_MSG) from exc
        return "\n\n".join(_docx_blocks(doc)), None
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    if _looks_binary(text):
        raise RuntimeError(UNSUPPORTED_MSG)
    return text, None


def _w(tag: str) -> str:
    return "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag


def _docx_blocks(doc: Any) -> List[str]:
    """Word 讲稿按文档里的顺序读出来：正文的段落一段一段；表格一行一段（同一行的格子、格子里的几段之间换行，
    每个格子算一句）。以前只读 doc.paragraphs，表格里的字全丢了（教案常常写在「教学环节 | 讲解内容」的表格里）。
    合并的格子只读一次；表格里套的表格、内容控件里的段落也读。文本框里的字还是读不到（见 docx_info）。"""
    from docx.text.paragraph import Paragraph  # type: ignore

    def para_text(p_el: Any) -> str:
        try:
            return Paragraph(p_el, doc).text.strip()
        except Exception:
            return "".join(t.text or "" for t in p_el.iter(_w("t"))).strip()

    def cell_lines(tc: Any) -> List[str]:
        out: List[str] = []
        for child in tc.iterchildren():
            if child.tag == _w("p"):
                t = para_text(child)
                if t:
                    out.append(t)
            elif child.tag == _w("tbl"):
                out += table_rows(child)
            elif child.tag == _w("sdt"):
                content = child.find(_w("sdtContent"))
                if content is not None:
                    out += cell_lines(content)
        return out

    def table_rows(tbl: Any) -> List[str]:
        rows: List[str] = []
        for tr in tbl.iterchildren(_w("tr")):
            parts: List[str] = []
            for tc in tr.iterchildren(_w("tc")):
                vmerge = tc.find(f"{_w('tcPr')}/{_w('vMerge')}")
                if vmerge is not None and vmerge.get(_w("val")) in (None, "continue"):
                    continue  # 竖着合并的格子：内容在最上面那一格，已经读过了
                parts += cell_lines(tc)
            if parts:
                rows.append("\n".join(parts))
        return rows

    blocks: List[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == _w("p"):
            t = para_text(child)
            if t:
                blocks.append(t)
        elif child.tag == _w("tbl"):
            blocks += table_rows(child)
        elif child.tag == _w("sdt"):
            content = child.find(_w("sdtContent"))
            if content is not None:
                blocks += cell_lines(content)
    return blocks


def docx_info(path: Path) -> Dict[str, int]:
    """Word 文件里有几个表格、几个文本框（上传讲稿时告诉老师：表格读进来了；文本框里的字读不到）。读不了时都是 0。"""
    try:
        import docx  # type: ignore

        body = docx.Document(str(path)).element.body
        return {"tables": len(body.findall(".//" + _w("tbl"))), "textboxes": len(body.findall(".//" + _w("txbxContent")))}
    except Exception:
        return {"tables": 0, "textboxes": 0}


# ----------------------------------------------------------------------------- 清理
def clean_markdown(text: str, skip_code_blocks: bool = True) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if skip_code_blocks:
        text = re.sub(r"```.*?```", "\n\n", text, flags=re.S)
    else:
        text = text.replace("```", "")
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)                    # 图片
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)               # 链接
    text = re.sub(r"<(?!break\b)[^>]+>", "", text)                      # HTML 标签（保留 <break>）
    lines = []
    for line in text.split("\n"):
        s = line.strip()
        if re.fullmatch(r"[-*_=|:\s]{3,}", s):                          # 分隔线 / 表格分隔
            lines.append("")
            continue
        heading = re.match(r"^#{1,6}\s+(.*)$", s)
        if heading:
            h = heading.group(1).strip()
            lines += ["", ensure_final_punct(h, detect_lang(h)), ""]
            continue
        s = re.sub(r"^>\s?", "", s)                                     # 引用
        s = re.sub(r"^(?:[-*+]|\d+[.)、])\s+", "", s)                   # 列表
        if s.startswith("|") and s.endswith("|"):                       # 表格行
            s = "，".join(c.strip() for c in s.strip("|").split("|") if c.strip())
        lines.append(s)
    text = "\n".join(lines)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"\2", text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = EMOJI_RE.sub("", text)
    return text


def _normalize_pause_marks(text: str) -> str:
    """【停顿=2】【停顿】［停顿＝２］→ [停顿=2] / [停顿]（统一以后认不出来的保持原样）。"""

    def fix(m: "re.Match[str]") -> str:
        cand = "[" + unicodedata.normalize("NFKC", m.group(1)).strip() + "]"
        return cand if PAUSE_RE.fullmatch(cand) else m.group(0)

    return _PAUSE_LOOSE_RE.sub(fix, text)


def apply_lexicon(text: str, lexicon: Sequence[Tuple[str, str]]) -> str:
    for src, dst in sorted(lexicon, key=lambda p: -len(p[0])):
        if re.fullmatch(r"[A-Za-z0-9 .+#-]+", src):
            text = re.sub(rf"(?<![A-Za-z]){re.escape(src)}(?![A-Za-z])", dst, text)
        else:
            text = text.replace(src, dst)
    return text


def tts_normalize(text: str) -> str:
    text = text.replace("——", "，").replace("--", "，")
    text = re.sub(r"(?<=\d)\s*[~～]\s*(?=\d)", "到" if count_cjk(text) else " to ", text)
    text = text.replace("~", " ").replace("～", " ")
    text = re.sub(r"[\[\]{}【】［］<>《》]", " ", text)  # 全角的［］也去掉（以前原样发给合成引擎）
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+(?=[，。！？；：,.!?;:])", "", text)
    return ensure_final_punct(text, detect_lang(text)) if text else text


# ----------------------------------------------------------------------------- 切句
def split_sentences(paragraph: str) -> List[str]:
    """按句末标点切句；英文句点要排除小数、缩写。"""
    out: List[str] = []
    buf = ""
    i, n = 0, len(paragraph)
    closers = "”\"'’）)」』】"
    while i < n:
        ch = paragraph[i]
        buf += ch
        boundary = False
        if ch in "。！？!?；;…":
            while i + 1 < n and (paragraph[i + 1] in "。！？!?…" or paragraph[i + 1] in closers):
                i += 1
                buf += paragraph[i]
            boundary = True
        elif ch == ".":
            prev_word = re.search(r"([A-Za-z.]+)\.$", buf)
            nxt = paragraph[i + 1] if i + 1 < n else ""
            is_decimal = bool(re.search(r"\d\.$", buf)) and nxt.isdigit()
            is_abbr = bool(prev_word) and prev_word.group(1).lower().rstrip(".") in ABBREVIATIONS
            single_initial = bool(re.search(r"(?:^|\s)[A-Z]\.$", buf))
            if not is_decimal and not is_abbr and not single_initial and (nxt == "" or nxt.isspace() or count_cjk(nxt)
                                                                           or nxt in closers):
                while i + 1 < n and paragraph[i + 1] in closers:
                    i += 1
                    buf += paragraph[i]
                boundary = True
        if boundary:
            if buf.strip():
                out.append(buf.strip())
            buf = ""
        i += 1
    if buf.strip():
        out.append(buf.strip())
    return out


def _split_clauses(sentence: str) -> List[str]:
    parts = re.split(f"(?<=[{re.escape(CLAUSE_CHARS)}])", sentence)
    return [p.strip() for p in parts if p.strip()]


#: 硬切时不能从中间切开的一段：英文单词 / 数字（连同里面的 ' . + # -）、一串空白、其它（中文和标点）
_HARD_TOKEN_RE = re.compile(r"[A-Za-z0-9'’.+#\-]+|\s+|[^A-Za-z0-9'’.+#\-\s]+")


def _hard_units(text: str) -> List[str]:
    """中文分句硬切用的最小单位：英文单词、数字整个算一个；中文有 jieba（整合包里有）时按词，没有时按字。"""
    from voicetwin.data.lexicon_fix import word_bounds

    units: List[str] = []
    for tok in _HARD_TOKEN_RE.findall(text):
        if tok.isspace() or not count_cjk(tok) or re.match(r"[A-Za-z0-9]", tok):
            units.append(tok)
            continue
        bounds = word_bounds(tok) if len(tok) > 1 else None
        if not bounds:
            units.extend(tok)
            continue
        cuts = sorted(b for b in bounds if 0 <= b <= len(tok))
        units.extend(tok[a:b] for a, b in zip(cuts, cuts[1:]) if b > a)
    return units


def _char_split(text: str, max_units: int) -> List[str]:
    """一个字一个字数，到上限就切（硬切的最后一道保险：保证每段都不超过上限）。原样切开，空格留在前一段末尾。"""
    out, cur = [], ""
    for ch in text:
        if cur.strip() and syllable_count(cur + ch) > max_units:
            out.append(cur)
            cur = ""
        cur += ch
    if cur.strip() or not out:
        out.append(cur)
    else:
        out[-1] += cur
    return out


def _is_latin(unit: str) -> bool:
    return bool(re.match(r"[A-Za-z0-9]", unit))


#: 硬切时每个单位（词、字、英文单词）的音节数：同样的词反复出现，记下来不用每次重算
_unit_syllables = lru_cache(maxsize=8192)(syllable_count)


def _cut_point(cur: List[str], nxt: str, limit: int) -> int:
    """cur 装满了、下一个单位是 nxt：挑切开的地方，返回 j（cur[:j] 成为前一段；空格留在前一段末尾，原样切开）。

    1. 不挨着英文单词的空格（用空格代替逗号的讲稿），前面已经有一半长：在最后一个这样的空格处切；
    2. 不然就在 nxt 前面切，可是不在两个英文单词（数字）中间切（以前会切成「which we，」「know……」，英文短语中间停一下）：
       往前找最后一个不在英文中间、前面也有一半长的地方；
    3. 都找不到：在 nxt 前面切。"""
    seq = cur + [nxt]

    def inside_en(j: int) -> bool:  # 切在 seq[j] 前面，是不是在两个英文单词（数字）中间
        a = next((x for x in reversed(seq[:j]) if not x.isspace()), "")
        b = next((x for x in seq[j:] if not x.isspace()), "")
        return _is_latin(a[-1:]) and _is_latin(b[:1])

    before = [0]  # before[j]：cur[:j] 有多少音节
    for x in cur:
        before.append(before[-1] + _unit_syllables(x))
    for j in range(len(cur) - 1, 0, -1):
        if (cur[j].isspace() and before[j] >= limit / 2 and not _is_latin(cur[j - 1][-1:])
                and not _is_latin(seq[j + 1][:1])):
            return j + 1
    if not inside_en(len(cur)):
        return len(cur)
    for j in range(len(cur) - 1, 0, -1):
        if not cur[j].isspace() and before[j] >= limit / 2 and not inside_en(j):
            return j
    return len(cur)


def _fill_pieces(units: List[str], limit: int, max_units: int) -> List[str]:
    """把单位一个一个装进段里，到 limit 就切；剩下的全部装得进一段（不超过 max_units）时不再切，
    免得最后剩下「错误」两个字（以前剩下的这一点会被 chunk_sentence 并回前一段，补的逗号就落在一段中间）。"""
    out: List[str] = []
    cur: List[str] = []
    cur_n = 0
    rest = sum(_unit_syllables(u) for u in units)  # 从这一个起、后面还有多少音节
    for u in units:
        n = _unit_syllables(u)
        if cur and cur_n + n > limit and cur_n + rest > max_units:
            at = _cut_point(cur, u, limit)
            out.append("".join(cur[:at]))
            cur = cur[at:]
            cur_n = sum(_unit_syllables(x) for x in cur)
            if cur and cur_n + n > limit and cur_n + rest > max_units:  # 切开处后面剩下的加上这一个还是太长：剩下的也单独成一段
                out.append("".join(cur))
                cur, cur_n = [], 0
        rest -= n
        if not cur and u.isspace() and out:
            out[-1] += u  # 段开头的空格留在前一段末尾
        else:
            cur.append(u)
            cur_n += n
    if cur:
        out.append("".join(cur))
    # 保险：英文单词和数字连着写、拆成字母以后音节的算法不一样……个别段还是超过上限时，这一段按字切（以前的办法）
    return [p for o in out if o for p in (_char_split(o, max_units) if syllable_count(o) > max_units else [o])]


def _hard_pieces(text: str, max_units: int) -> List[str]:
    """硬切成几段，每段都不超过上限；原样切开（"".join(结果) == text，空格留在前一段末尾），不补逗号。

    逗号由 chunk_sentence 在真的切开的地方补（_hard_split 是补好逗号的版本）：chunk_sentence 会把装得下的几段并成一段，
    以前每段先补好「，」，并回去以后逗号就落在一段中间（「翻译成一个，形容词。」），合成时在那里停一下、字幕里也看得到。"""
    if count_cjk(text) == 0:
        out, cur = [], ""
        for w in re.findall(r"\S+\s*|\s+", text):
            if cur.strip() and syllable_count(cur + w) > max_units:
                out.append(cur)
                cur = ""
            cur += w
        if cur:
            out.append(cur)
        # 一个「单词」本身就超过上限（很长的一串数字、网址）：只好按字切
        return [p for o in out for p in (_char_split(o, max_units) if syllable_count(o) > max_units else [o])]
    # 中文分句切成长短差不多的几段（以前每段塞满 50 个字，最后剩下「东西。」两个字单独合成）
    total = syllable_count(text)
    k = max(1, math.ceil(total / max(1, max_units)))
    hard = _hard_units(text)
    best: Optional[List[str]] = None
    for kk in range(k, k + 4):
        limit = max(1, math.ceil(total / kk))
        units: List[str] = []
        for u in hard:
            # 一个单位本身就太长（很长的英文单词、jieba 认成一个词的长串）：只好按字切
            units.extend(list(u) if _unit_syllables(u) > limit else [u])
        out = _fill_pieces(units, limit, max_units)
        if len(out) <= kk:  # 按计划切成了 kk 段
            return out
        # 词的边界凑不出 kk 段（比如 49 + 49 + 剩下 2 个字）：多切一段再试，段数最少的留着
        if best is None or len(out) < len(best):
            best = out
    return best or [text]


def _add_cut_comma(text: str, piece: str) -> str:
    """硬切开的地方补一个逗号（不补句号：以前补「。」，一个词被切成两半「特。」「别注意」，语调也落下来）。
    piece 是切开处前面那一截：有汉字补「，」，全是英文补「,」。"""
    text = text.rstrip()
    if text and text[-1] not in CLAUSE_CHARS + SENT_END_CHARS + ".,":
        text += "，" if count_cjk(piece) else ","
    return text


def _hard_split(text: str, max_units: int) -> List[str]:
    """一个分句仍然太长、中间又没有逗号（很少见，多半是从语音转文字工具里复制来的、用空格代替逗号的讲稿）：硬切。

    英文按单词切；中文切成长短差不多的几段：优先在不挨着英文的空格处切（空格前面已经有一半长），不然在词和词之间切
    （有 jieba 时），英文单词、数字不切开，也尽量不从两个英文单词中间切开；最后剩下的装得进一段就不再切。
    除了最后一段，每段末尾补一个逗号。（chunk_sentence 用的是不补逗号的 _hard_pieces，只在真的切开的地方补逗号。）"""
    out = [p.strip() for p in _hard_pieces(text, max_units)]
    out = [p for p in out if p]
    return [_add_cut_comma(p, p) if i < len(out) - 1 else p for i, p in enumerate(out)]


def chunk_sentence(sentence: str, max_units: int) -> List[str]:
    if syllable_count(sentence) <= max_units:
        return [sentence]
    chunks: List[str] = []
    cur, cut_tail = "", ""  # cut_tail：cur 的最后一截是硬切出来的、同一个分句后面还有（这里结束一段就要补逗号）
    for clause in _split_clauses(sentence):
        pieces = [clause] if syllable_count(clause) <= max_units else [p for p in _hard_pieces(clause, max_units)
                                                                       if p.strip()]
        for k, piece in enumerate(pieces):
            if cut_tail:  # 同一个分句硬切开的下一截：原样接上（没有在这里切，就不补逗号）
                joiner = ""
            else:
                joiner = "" if (count_cjk(cur[-1:]) or count_cjk(piece[:1]) or not cur) else " "
            if cur and syllable_count(cur + joiner + piece) > max_units:
                chunks.append(_add_cut_comma(cur, cut_tail) if cut_tail else cur)
                cur = piece.lstrip()
            else:
                cur = cur + joiner + piece if cur else piece
            cut_tail = piece if k < len(pieces) - 1 else ""
    if cur:
        chunks.append(cur.rstrip())
    return chunks


def _join(a: str, b: str) -> str:
    return a + ("" if count_cjk(a[-1:]) or count_cjk(b[:1]) else " ") + b


def _is_script_file(source: str) -> bool:
    """是讲稿文件的路径（不是讲稿文字）。问硬盘出错（一行太长等）就当文字。"""
    if len(source) >= 1024 or "\n" in source:
        return False
    try:
        return Path(source).suffix.lower() in SCRIPT_EXTS and Path(source).exists()
    except (OSError, ValueError):
        return False


# ----------------------------------------------------------------------------- 主函数
def parse_script(source: Union[str, Path], lexicon: Sequence[Tuple[str, str]] = (), max_units_zh: int = 50,
                 max_units_en: int = 45, min_units: int = 6, skip_code_blocks: bool = True) -> List[ScriptSegment]:
    cues = None
    if isinstance(source, Path) or (isinstance(source, str) and _is_script_file(source)):
        text, cues = read_script_file(Path(source))
    else:
        text = str(source)
    if cues is not None:
        return _segments_from_cues(cues, lexicon, max_units_zh, max_units_en)

    text = _normalize_pause_marks(clean_markdown(text, skip_code_blocks))
    # 把停顿标记替换成独立占位行
    tokens: List[Union[str, float]] = []
    pos = 0
    for m in PAUSE_RE.finditer(text):
        tokens.append(text[pos:m.start()])
        num, unit = (m.group(1), m.group(2)) if m.group(1) or m.group(0).startswith("[") else (m.group(3), m.group(4))
        if num:
            u = (unit or "").lower()
            sec = float(num) / 1000.0 if u in ("ms", "毫秒") else float(num) * (60.0 if u in ("分钟", "min") else 1.0)
            tokens.append(sec)
        else:
            tokens.append(-1.0)  # 段落停顿
        pos = m.end()
    tokens.append(text[pos:])

    segments: List[ScriptSegment] = []
    para_idx = 0
    lead = 0.0  # 写在第一句前面的停顿秒数（开头多留的静音）
    for tok in tokens:
        if isinstance(tok, float):
            if segments:
                segments[-1].pause_after = "paragraph" if tok < 0 else tok
            elif tok >= 0:
                lead = tok  # 讲稿最前面的 [停顿=3]：以前直接丢掉了，开头还是只有 0.35 秒
            continue
        paragraphs = re.split(r"\n\s*\n", tok)
        for p_i, para in enumerate(paragraphs):
            lines = [ln.strip() for ln in para.split("\n") if ln.strip()]
            if not lines:
                continue
            # 合并被硬换行打断的英文句子；其它换行视为句子结束
            merged = lines[0]
            for ln in lines[1:]:
                if re.search(r"[A-Za-z,]$", merged) and re.match(r"^[a-z]", ln):
                    merged += " " + ln
                else:
                    merged = ensure_final_punct(merged, detect_lang(merged)) + ("" if count_cjk(ln[:1]) else " ") + ln
            sentences = split_sentences(merged)
            # 太短的句子并入下一句（例如"好。"单独合成不自然）
            packed: List[str] = []
            for s in sentences:
                limit = max_units_zh if detect_lang(s) == "zh" else max_units_en
                if packed and syllable_count(packed[-1]) < min_units and syllable_count(_join(packed[-1], s)) <= limit:
                    packed[-1] = _join(packed[-1], s)
                else:
                    packed.append(s)
            if len(packed) > 1 and syllable_count(packed[-1]) < min_units:
                last = packed.pop()
                packed[-1] = _join(packed[-1], last)
            for s_i, sent in enumerate(packed):
                lang = detect_lang(sent)
                chunks = chunk_sentence(sent, max_units_zh if lang == "zh" else max_units_en)
                for c_i, chunk in enumerate(chunks):
                    last_chunk = c_i == len(chunks) - 1
                    pause: Union[str, float] = "clause" if not last_chunk else "sentence"
                    if last_chunk and s_i == len(packed) - 1:
                        pause = "paragraph"
                    display = chunk.strip()
                    tts_text = tts_normalize(apply_lexicon(display, lexicon))
                    if not syllable_count(tts_text):
                        continue
                    segments.append(ScriptSegment(text=tts_text, display=display, lang=detect_lang(tts_text),
                                                  kind=sentence_kind(sent), pause_after=pause, paragraph=para_idx))
            para_idx += 1
    if segments and segments[-1].pause_after == "paragraph":
        segments[-1].pause_after = 0.0
    if segments and lead > 0:
        segments[0].extra["pause_before"] = lead  # 拼接时第一句前面留这么久（engine._layout）
    for i, s in enumerate(segments):
        s.index = i
        bad = _UNREAD_MARK_RE.search(s.display)
        if bad:
            s.extra["unread_mark"] = bad.group(0)  # 没认出来的停顿标记：会被读出来，生成时提醒老师
    return segments


def _segments_from_cues(cues: list, lexicon, max_units_zh: int, max_units_en: int) -> List[ScriptSegment]:
    segments: List[ScriptSegment] = []
    for c_i, cue in enumerate(cues):
        lang = detect_lang(cue.text)
        chunks = chunk_sentence(cue.text, max_units_zh if lang == "zh" else max_units_en)
        for k, chunk in enumerate(chunks):
            tts_text = tts_normalize(apply_lexicon(chunk, lexicon))
            if not syllable_count(tts_text):
                continue
            segments.append(ScriptSegment(
                text=tts_text, display=chunk, lang=detect_lang(tts_text), kind=sentence_kind(cue.text),
                pause_after="clause" if k < len(chunks) - 1 else 0.0, paragraph=c_i,
                cue_start=cue.start if k == 0 else None, cue_end=cue.end if k == len(chunks) - 1 else None,
            ))
    for i, s in enumerate(segments):
        s.index = i
    return segments
