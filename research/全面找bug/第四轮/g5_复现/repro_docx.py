import sys
import os; sys.path.insert(0, os.environ.get("VT_ROOT", str(__import__("pathlib").Path(__file__).resolve().parents[4])))
import docx
from pathlib import Path
from voicetwin.synth.script import read_script_file, parse_script, docx_info
d = docx.Document()
d.add_paragraph("第三课 定语从句")
t = d.add_table(rows=3, cols=2)
t.cell(0, 0).text = "教学环节"; t.cell(0, 1).text = "讲解内容"
t.cell(1, 0).text = "导入"; t.cell(1, 1).text = "同学们好，今天我们学习定语从句。"
t.cell(2, 0).text = "讲解"; t.cell(2, 1).text = "先看一个例句：The book which I bought is new."
m = t.cell(1, 0).merge(t.cell(2, 0))  # 竖着合并
d.add_paragraph("下课。")
import tempfile
out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(tempfile.mkdtemp()) / "教案.docx")
d.save(str(out))
text, _ = read_script_file(out)
print(repr(text))
print(docx_info(out))
for s in parse_script(text):
    print(repr(s.text), s.pause_after)
