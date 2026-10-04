"""第二轮的标黄表格。用法：python3 xlsx2.py 老师的transcripts.csv 输出目录"""
import csv, io, json, difflib, sys
sys.path.insert(0, '.')
from spec2 import UNSURE, INTRO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

SRC, OUT = sys.argv[1], sys.argv[2]
rows = list(csv.reader(io.StringIO(open(SRC, encoding='utf-8-sig', newline='').read(), newline='')))
head, data = rows[0], rows[1:]
TI, KI, DI = head.index('text'), head.index('keep'), head.index('drop_reason')
res = json.load(open(f'{OUT}/result.json'))
new = {int(k): v for k, v in res['new'].items()}
uns = {}
for no, a, g, why in UNSURE:
    uns.setdefault(no, []).append((a, g, why))

YELLOW = PatternFill('solid', fgColor='FFFF00')
BLUE = PatternFill('solid', fgColor='BDD7EE')
HEAD = PatternFill('solid', fgColor='D9D9D9')
RED = InlineFont(rFont='微软雅黑', sz=12, b=True, color='C00000')
NORM = InlineFont(rFont='微软雅黑', sz=12)
WRAP = Alignment(wrap_text=True, vertical='top')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F = Font(name='微软雅黑', size=12)
FB = Font(name='微软雅黑', size=12, bold=True)

def rich(text, spans):
    flags = [False] * len(text)
    for s, e in spans:
        for i in range(s, e):
            flags[i] = True
    out, i = [], 0
    while i < len(text):
        j = i
        while j < len(text) and flags[j] == flags[i]:
            j += 1
        out.append(TextBlock(RED if flags[i] else NORM, text[i:j]))
        i = j
    return CellRichText(out) if any(flags) else text

def changed(old, nt):
    sn, so = [], []
    for op, a1, a2, b1, b2 in difflib.SequenceMatcher(None, old, nt, autojunk=False).get_opcodes():
        if op == 'equal':
            continue
        if b2 > b1: sn.append((b1, b2))
        if a2 > a1: so.append((a1, a2))
        if b2 == b1 and b1 > 0: sn.append((b1 - 1, b1))
    return sn, so

def sheet(ws, header, widths):
    ws.append(header)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = FB, HEAD, WRAP, BOX
    ws.freeze_panes = 'A2'

def style(ws, r):
    for c in ws[r]:
        c.font, c.alignment, c.border = F, WRAP, BOX

wb = Workbook()
ws = wb.active; ws.title = '先看这里'
ws.column_dimensions['A'].width = 110
for t, bold, fill in INTRO(len(data), len(new), len(UNSURE), len(res['pairs'])):
    ws.append([t]); c = ws.cell(ws.max_row, 1); c.font = FB if bold else F; c.alignment = WRAP
    if fill == 'Y': c.fill = YELLOW
    if fill == 'B': c.fill = BLUE

if new:
    ws = wb.create_sheet('这次又改的%d句' % len(new))
    sheet(ws, ['行号', 'id', '纠正后的文字（红色是改过的字）', '你发来的文字（红色是被改掉的字）', '改了什么、为什么'], [7, 30, 60, 60, 45])
    for no in sorted(new):
        old, (nt, why) = data[no - 1][TI], new[no]
        sn, so = changed(old, nt)
        ws.append([no, data[no - 1][0], rich(nt, sn), rich(old, so), '；'.join(why)])
        style(ws, ws.max_row); ws.cell(ws.max_row, 3).fill = YELLOW

if uns:
    ws = wb.create_sheet('没把握（请听录音）')
    sheet(ws, ['行号', 'id', '现在的文字（红色是没把握的地方）', '可能是', '为什么觉得可能不对', '听了以后'], [7, 30, 60, 40, 50, 14])
    for no in sorted(uns):
        t = new.get(no, [data[no - 1][TI]])[0]
        for a, g, why in uns[no]:
            s = t.index(a)
            ws.append([no, data[no - 1][0], rich(t, [(s, s + len(a))]), g, why, ''])
            style(ws, ws.max_row); ws.cell(ws.max_row, 3).fill = BLUE

ws = wb.create_sheet('全部句子')
sheet(ws, ['行号', 'id', '用不用', '文字（纠正后）', '说明'], [7, 30, 12, 80, 50])
for no, r in enumerate(data, 1):
    old = r[TI]
    keep = '用' if r[KI] == '1' else '不用（%s）' % (r[DI] or '')
    if no in new:
        nt, why = new[no]
        ws.append([no, r[0], keep, rich(nt, changed(old, nt)[0]), '这次又改：' + '；'.join(why)])
        style(ws, ws.max_row); ws.cell(ws.max_row, 4).fill = YELLOW
        if no in uns:
            ws.cell(ws.max_row, 5).value += '｜还有没把握的：' + '；'.join(g for a, g, w in uns[no]) + '（请听录音）'
    elif no in uns:
        ws.append([no, r[0], keep, old, '没改，请听录音：可能是 ' + '；'.join(g for a, g, w in uns[no])])
        style(ws, ws.max_row); ws.cell(ws.max_row, 4).fill = BLUE
    else:
        ws.append([no, r[0], keep, old, ''])
        style(ws, ws.max_row)

if res['pairs']:
    ws = wb.create_sheet('在程序里怎么改')
    ws.append(['在声音分身的「校对表」上面：在「查找」框里输入左边的字 → 「替换成」框里输入右边的字 → 点「全部替换」。全部做完后点「保存修改」，再点「✅ 确认训练素材」。'])
    ws.merge_cells('A1:E1'); ws['A1'].font = FB; ws['A1'].alignment = WRAP; ws.row_dimensions[1].height = 48
    ws.append(['第几条', '行号', '查找（复制这个）', '替换成（复制这个）', '做完打勾'])
    for c in ws[2]:
        c.font, c.fill, c.alignment, c.border = FB, HEAD, WRAP, BOX
    for i, w in enumerate([8, 7, 45, 45, 10], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A3'
    for i, (no, rid, f, t) in enumerate(res['pairs'], 1):
        ws.append([i, no, f, t, ''])
        style(ws, ws.max_row); ws.cell(ws.max_row, 4).fill = YELLOW

wb.save(f'{OUT}/第二次校对结果_标黄.xlsx')
print('ok')
