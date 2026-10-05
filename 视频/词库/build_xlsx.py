import json,sys
from openpyxl import Workbook
from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
from openpyxl.utils import get_column_letter as L
arts=json.load(open('rev_01_35.json'))+json.load(open('rev_36_70.json'))
arts.sort(key=lambda a:a['no'])
wb=Workbook(); ws=wb.active; ws.title='70篇重点词汇'
hf=PatternFill('solid',fgColor='1E3A2F'); sf=PatternFill('solid',fgColor='DCEFE5')
thin=Side(style='thin',color='B7CFC2'); B=Border(left=thin,right=thin,top=thin,bottom=thin)
for k,a in enumerate(arts):
    c=2*k+1
    ws.merge_cells(start_row=1,start_column=c,end_row=1,end_column=c+1)
    h=ws.cell(1,c,f"{a['no']} {a['title']}"); h.font=Font(bold=True,color='FFFFFF'); h.fill=hf
    h.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for j,t in enumerate(['英文','中文']):
        x=ws.cell(2,c+j,t); x.font=Font(bold=True); x.fill=sf; x.alignment=Alignment(horizontal='center'); x.border=B
    for r,(en,zh) in enumerate(a['items']):
        e=ws.cell(3+r,c,en); z=ws.cell(3+r,c+1,zh); e.border=z.border=B
        e.font=Font(name='Calibri',size=11); z.font=Font(name='微软雅黑',size=11)
    ws.column_dimensions[L(c)].width=22; ws.column_dimensions[L(c+1)].width=20
ws.row_dimensions[1].height=36; ws.freeze_panes='A3'
wb.save('70篇重点词汇.xlsx'); print(sum(len(a['items']) for a in arts),len(arts))
