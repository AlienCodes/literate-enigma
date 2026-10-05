import json,re,sys
from docx import Document
from docx.shared import Pt,RGBColor,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import render as R
PAL=[(251,191,36),(56,189,248),(244,114,182),(190,242,100),(196,181,253),(251,146,60),(94,234,212),(252,165,165)]
BG='0E241E'; EN=(240,247,242); ZH=(170,196,182)
d=json.load(open(sys.argv[1]))
colors={};i=0
for s in d['sentences']:
    for c in s['chunks']:
        for w in re.findall(r'\*\*([^*]+)\*\*',c['en']):
            if w.lower() not in colors: colors[w.lower()]=PAL[i%8];i+=1
R.unify_colors(colors,d['sentences'])
doc=Document()
sec=doc.sections[0]; sec.left_margin=sec.right_margin=Cm(1.5)
bg=OxmlElement('w:background'); bg.set(qn('w:color'),BG); doc.element.insert(0,bg)
st=OxmlElement('w:displayBackgroundShape'); se=doc.settings.element
z=se.find(qn('w:zoom'))
(z.addnext(st) if z is not None else se.insert(0,st))
def run(p,t,col,size,bold=False,zh=False):
    r=p.add_run(t); r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor(*col)
    r.font.name='Calibri'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
def para(t,col,size,bold=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,t,col,size,bold); return p
para(d['no'],(110,140,125),10); para(d['title_en'],EN,20,True); para(d['title_zh'],ZH,13)
def noborder(t):
    tbl=t._tbl; pr=tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'nil'); b.append(x)
    pr.append(b)
prev=None
for s in d['sentences']:
    if prev is not None and s['para']!=prev: doc.add_paragraph()
    prev=s['para']
    for c in s['chunks']:
      segs=[[]]
      for e,z in c['align']:
        if e=='\n': segs.append([])
        else: segs[-1].append((e,z))
      for al in segs:
        t=doc.add_table(rows=2,cols=len(al)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True; noborder(t)
        for j,(e,z) in enumerate(al):
            for row,lines,size,dc in ((0,[R._segs_en(e,colors)],13,EN),(1,[R._segs_zh(x,colors) for x in z.split('\n')],10.5,ZH)):
                p=t.cell(row,j).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0)
                for k,ss in enumerate(lines):
                    if k: p.add_run().add_break()  # 中文格内换行
                    for tx,col in ss: run(p,tx,col or dc,size,bool(col))
      if c.get('note'): para(c['note'].replace('#Y',''),(251,191,36) if c['note'].startswith('#Y') else ZH,10)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
doc.save(sys.argv[2])
